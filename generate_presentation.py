"""
Script to generate the official executive PowerPoint Presentation (.pptx) for AnnaSetu AI.
Formatted specifically for executive presentation and technical project walkthrough.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck(filename="AnnaSetu_AI_Executive_Presentation.pptx"):
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Colors
    c_primary = RGBColor(6, 78, 59)      # Deep Forest Emerald #064e3b
    c_secondary = RGBColor(4, 120, 87)   # Emerald Green #047857
    c_accent = RGBColor(217, 119, 6)     # Warm Amber #d97706
    c_dark = RGBColor(15, 23, 42)        # Slate 900 #0f172a
    c_body = RGBColor(51, 65, 85)        # Slate 700 #334155
    c_card_bg = RGBColor(248, 250, 252)  # Slate 50 #f8fafc
    c_card_border = RGBColor(203, 213, 225) # Slate 300 #cbd5e1
    c_white = RGBColor(255, 255, 255)
    c_rose = RGBColor(159, 18, 57)       # Wine / Rose #9f1239
    c_purple = RGBColor(107, 33, 168)    # Purple #6b21a8

    def add_header(slide, title_text, category_text="UN SDG 2 & 12 | ANNASETU AI"):
        # Category / breadcrumb
        tb_cat = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        tf_cat = tb_cat.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = c_accent

        # Title
        tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.6))
        tf_title = tb_title.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = c_primary

        # Underline accent
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.733), Inches(0.03))
        line.fill.solid()
        line.fill.fore_color.rgb = c_card_border
        line.line.color.rgb = c_card_border

    def add_card(slide, left, top, width, height, title, body_bullets, title_color=c_primary, bg_color=c_card_bg):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = c_card_border
        card.line.width = Pt(1)

        tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), width - Inches(0.4), height - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True

        if title:
            p_t = tf.paragraphs[0]
            p_t.text = title
            p_t.font.size = Pt(14)
            p_t.font.bold = True
            p_t.font.color.rgb = title_color
            p_t.space_after = Pt(8)

        for i, b in enumerate(body_bullets):
            p = tf.add_paragraph() if (title or i > 0) else tf.paragraphs[0]
            p.text = f"•  {b}"
            p.font.size = Pt(11)
            p.font.color.rgb = c_body
            p.space_after = Pt(5)

    # ==========================================
    # SLIDE 1: TITLE SLIDE
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    # Background full color banner
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = c_primary
    bg1.line.fill.background()

    # Content box
    tb_title = s1.shapes.add_textbox(Inches(1.2), Inches(1.5), Inches(11.0), Inches(4.5))
    tf1 = tb_title.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "SMART SURPLUS FOOD REDISTRIBUTION PLATFORM • UN SDG 2 & 12"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = c_accent
    p0.space_after = Pt(12)

    p1 = tf1.add_paragraph()
    p1.text = "AnnaSetu AI"
    p1.font.size = Pt(46)
    p1.font.bold = True
    p1.font.color.rgb = c_white
    p1.space_after = Pt(6)

    p2 = tf1.add_paragraph()
    p2.text = "Smart Surplus Food Redistribution Platform"
    p2.font.size = Pt(24)
    p2.font.color.rgb = RGBColor(167, 243, 208) # Mint light
    p2.space_after = Pt(18)

    p3 = tf1.add_paragraph()
    p3.text = "“Share More. Waste Less. Feed Hope.”"
    p3.font.size = Pt(18)
    p3.font.italic = True
    p3.font.color.rgb = c_white
    p3.space_after = Pt(28)

    p4 = tf1.add_paragraph()
    p4.text = "Global Goals Alignment: SDG 2 (Zero Hunger) & SDG 12 (Responsible Consumption & Production)"
    p4.font.size = Pt(13)
    p4.font.bold = True
    p4.font.color.rgb = RGBColor(254, 243, 199) # Warm yellow

    # ==========================================
    # SLIDE 2: PROBLEM STATEMENT & URGENCY
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "The Challenge: Urban Food Waste vs. Hunger Paradox")

    add_card(s2, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2),
             "The Reality in Urban India",
             [
                 "50+ Million Tonnes of edible food discarded in Indian urban centers annually.",
                 "Commercial Banquets, Hotels, and College Hostels routinely prepare 15–25% surplus that gets dumped.",
                 "Landfill Food Decomposition emits potent Methane (CH4) greenhouse gas, exacerbating climate degradation.",
                 "Meanwhile, millions in urban shelters, child relief centers, and slums face daily nutritional deficits."
             ], title_color=c_rose)

    add_card(s2, Inches(6.8), Inches(1.7), Inches(5.6), Inches(5.2),
             "Why Existing Methods Fail",
             [
                 "Lack of Digital Coordination: Kitchens rely on ad-hoc phone calls during late-night closing hours.",
                 "Food Safety Liabilities: Donors fear litigation without strict temperature holding parameters & legal declarations.",
                 "Mismatched Batch Sizes: Large banquets overwhelm small local shelters, while large community kitchens get insufficient batches.",
                 "Logistical Bottlenecks: Absence of organized, verified volunteer courier networks with insulated thermal carriers."
             ], title_color=c_accent)

    # ==========================================
    # SLIDE 3: THE ANNASETU AI SOLUTION
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "The Solution: AnnaSetu AI Platform Architecture")

    add_card(s3, Inches(0.8), Inches(1.7), Inches(3.6), Inches(5.2),
             "1. Food Donors",
             [
                 "Post surplus batches in seconds.",
                 "AI Smart Suggestions auto-fill category, detect allergens, and enhance packaging notes.",
                 "Review proximity match recommendations with transparent scores.",
                 "Approve recipient reservation requests.",
                 "Monitor real-time pickup transit."
             ], title_color=c_primary)

    add_card(s3, Inches(4.8), Inches(1.7), Inches(3.6), Inches(5.2),
             "2. Recipient Charities (NGOs)",
             [
                 "Browse catalog within service radius.",
                 "AI Natural Language Query parser (e.g. '50 veg meals for dinner in Indiranagar').",
                 "Reserve batches capped by daily shelter distribution capacity.",
                 "Accreditation Security Gate: Unverified NGOs locked until 12A/80G/Darpan approval."
             ], title_color=c_secondary)

    add_card(s3, Inches(8.8), Inches(1.7), Inches(3.6), Inches(5.2),
             "3. Volunteer Couriers",
             [
                 "Browse city-wide unassigned runs.",
                 "Claim runs & access route details.",
                 "4-Stage Delivery State Machine: ASSIGNED → ACCEPTED → PICKED_UP → COMPLETED.",
                 "FSSAI Thermal Carrier Pledge: hot (≥60°C) and chilled (≤5°C) compliance."
             ], title_color=c_purple)

    # ==========================================
    # SLIDE 4: UN SUSTAINABLE DEVELOPMENT GOALS
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Alignment with United Nations Sustainable Development Goals")

    add_card(s4, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2),
             "SDG 2: Zero Hunger (Target 2.1 & 2.2)",
             [
                 "Direct Channel to Verified Shelters: Wholesome meals are redirected from hotel banquets directly into community dining halls.",
                 "Capacity-Aware Allocation: Batch sizes match recipient intake capability, guaranteeing food is consumed rather than re-wasted.",
                 "Allergen & Dietary Safeguards: Protects vulnerable beneficiaries through explicit disclosure of nuts, dairy, gluten, and vegan/jain attributes.",
                 "Real-Time Plate Tracking: Quantifies exact nutritional impact through SQL-aggregated servings count."
             ], title_color=c_accent)

    add_card(s4, Inches(6.8), Inches(1.7), Inches(5.6), Inches(5.2),
             "SDG 12: Responsible Consumption (Target 12.3)",
             [
                 "Commercial Kitchen Waste Reduction: Establishes institutional discipline in hotels, hostels, and caterers.",
                 "Landfill Methane Prevention: Edible organic matter diverted from rotting in open municipal dump sites.",
                 "Carbon & Kilogram Measurement: Direct SQL computation of total kilograms diverted from municipal waste streams.",
                 "Circular Urban Food Economy: Demonstrates a replicable blueprint for smart, zero-waste metropolitan centers."
             ], title_color=c_primary)

    # ==========================================
    # SLIDE 5: FULL-STACK TECHNICAL ARCHITECTURE
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Production-Grade Full-Stack Technology Stack")

    # Table of Tech Stack
    t_shape = s5.shapes.add_table(5, 3, Inches(0.8), Inches(1.7), Inches(11.7), Inches(5.0))
    table = t_shape.table
    table.columns[0].width = Inches(2.2)
    table.columns[1].width = Inches(3.5)
    table.columns[2].width = Inches(6.0)

    headers = ["Layer", "Technology Selection", "Engineering Rationale & Capabilities"]
    for col_idx, h in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = c_primary
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = c_white

    rows = [
        ("Frontend Client", "React 19 + TypeScript + Vite + Tailwind CSS v4", "High-performance reactive UI with type safety, accessible design, and Recharts impact analytics."),
        ("Backend Framework", "Python 3.12/3.14 + FastAPI + Pydantic v2", "Asynchronous high-throughput REST APIs, strict schema validation, automated OpenAPI docs (/docs)."),
        ("Database & ORM", "SQLAlchemy 2.0 ORM + Alembic + SQLite/Postgres", "Relational data integrity, schema migration tracking, and row-level locking (SELECT FOR UPDATE)."),
        ("Security & AI", "Argon2 + HS256 JWT + Google Gemini Python SDK", "Cryptographic password hashing, stateless token authorization, and resilient AI with deterministic fallbacks.")
    ]

    for row_idx, r_data in enumerate(rows):
        for col_idx, val in enumerate(r_data):
            cell = table.cell(row_idx + 1, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = c_white if row_idx % 2 == 0 else c_card_bg
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(10)
            p.font.color.rgb = c_dark if col_idx == 0 else c_body
            if col_idx == 0:
                p.font.bold = True

    # ==========================================
    # SLIDE 6: CORE TECHNICAL INNOVATIONS
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Core Technical Innovations & Engineering Highlights")

    add_card(s6, Inches(0.8), Inches(1.7), Inches(5.6), Inches(2.5),
             "1. FSSAI Conservative Safety Protocol",
             [
                 "Safety rules strictly override raw AI predictions.",
                 "Enforces hot-holding (≥60°C) and refrigerated (≤5°C) constraints.",
                 "Mandatory legal food safety declarations required before posting."
             ], title_color=c_primary)

    add_card(s6, Inches(6.8), Inches(1.7), Inches(5.6), Inches(2.5),
             "2. Transparent Proximity Matching (0–100)",
             [
                 "Haversine Geodesic Distance Formula (≤ 40 points).",
                 "Shelter Daily Distribution Capacity Fit (≤ 30 points).",
                 "Schedule Urgency Window (≤ 20 pts) & Dietary Match (≤ 10 pts)."
             ], title_color=c_accent)

    add_card(s6, Inches(0.8), Inches(4.4), Inches(5.6), Inches(2.5),
             "3. Concurrency & Atomic Inventory Locking",
             [
                 "Database row-level transactions (SELECT FOR UPDATE).",
                 "Eliminates race conditions and over-allocation across multiple NGOs.",
                 "Zero double-booking guarantee under concurrent requests."
             ], title_color=c_rose)

    add_card(s6, Inches(6.8), Inches(4.4), Inches(5.6), Inches(2.5),
             "4. 4-Stage Volunteer Delivery State Machine",
             [
                 "Auditable transitions: ASSIGNED → ACCEPTED → PICKED_UP → COMPLETED.",
                 "Thermal container hygiene validation and courier transit notes.",
                 "Real-time notifications dispatched to all participants."
             ], title_color=c_purple)

    # ==========================================
    # SLIDE 7: RESILIENT GOOGLE GEMINI AI
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Google Gemini AI Co-Pilot with Deterministic Fallbacks")

    add_card(s7, Inches(0.8), Inches(1.7), Inches(5.6), Inches(2.5),
             "A. Auto-Categorization & Allergen Detection",
             [
                 "Reads meal titles and descriptions to classify category.",
                 "Detects allergens (Dairy, Gluten, Peanuts, Eggs, Soy).",
                 "Recommends thermal holding protocol (hot_holding / refrigerated)."
             ], title_color=c_primary)

    add_card(s7, Inches(6.8), Inches(1.7), Inches(5.6), Inches(2.5),
             "B. Listing Quality Enhancement Co-Pilot",
             [
                 "Computes listing completeness score (0 to 100).",
                 "Generates structured packaging and hygiene improvement advice.",
                 "Prompts donors for missing storage or handling details."
             ], title_color=c_secondary)

    add_card(s7, Inches(0.8), Inches(4.4), Inches(5.6), Inches(2.5),
             "C. Natural Language Search Parser",
             [
                 "Translates conversational queries into structured SQL filters.",
                 "e.g. 'Need 50 veg meals for dinner in Indiranagar' → category='prepared_meals', dietary='vegetarian', urgency=True."
             ], title_color=c_accent)

    add_card(s7, Inches(6.8), Inches(4.4), Inches(5.6), Inches(2.5),
             "D. 100% Deterministic Resilient Fallback",
             [
                 "Zero single-point-of-failure: if Gemini API is unset or throttled, built-in rule engines take over instantly.",
                 "Full functionality maintained offline or without cloud API keys."
             ], title_color=c_purple)

    # ==========================================
    # SLIDE 8: VERIFICATION & TESTING RESULTS
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Quality Assurance: 100% Automated Test Suite Success")

    add_card(s8, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2),
             "Backend Automated Pytest Suite",
             [
                 "17 / 17 Tests Passing (100% Green).",
                 "Authentication & Cryptography: Argon2 hashing, JWT validity.",
                 "RBAC Security: Strict role isolation across all protected APIs.",
                 "Accreditation Gate: Unverified NGOs blocked from live requests.",
                 "Input Validation: Past deadlines and negative quantities rejected.",
                 "Concurrency: Atomic inventory locking prevents double-booking.",
                 "State Machine: Couriers restricted to valid status transitions.",
                 "Live E2E Test: Validates complete donor-to-beneficiary lifecycle."
             ], title_color=c_secondary)

    add_card(s8, Inches(6.8), Inches(1.7), Inches(5.6), Inches(5.2),
             "Frontend & Code Quality",
             [
                 "TypeScript Strict Mode: 0 Compilation Errors.",
                 "Production Build: Verified with npm run build (1.53s build time).",
                 "UI / UX Design Standards: Responsive across desktop, tablet, and mobile.",
                 "Data Visualization: Live Recharts graphs powered by database SQL.",
                 "Tamper-Evident Audit Ledger: Append-only ledger recording all events.",
                 "Security Best Practices: GEMINI_API_KEY never exposed in frontend bundle."
             ], title_color=c_primary)

    # ==========================================
    # SLIDE 9: LIVE DEMONSTRATION & CREDENTIALS
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Platform Demonstration & Pre-Seeded Accounts")

    # Table of Demo Accounts
    t_shape9 = s9.shapes.add_table(5, 4, Inches(0.8), Inches(1.7), Inches(11.7), Inches(5.0))
    table9 = t_shape9.table
    table9.columns[0].width = Inches(2.0)
    table9.columns[1].width = Inches(2.8)
    table9.columns[2].width = Inches(1.8)
    table9.columns[3].width = Inches(5.1)

    headers9 = ["Persona / Role", "Login Email", "Password", "Evaluation Walkthrough Focus"]
    for col_idx, h in enumerate(headers9):
        cell = table9.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = c_primary
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = c_white

    rows9 = [
        ("Central Administrator", "admin@annasetu.ai", "Admin@123456", "Review NGO verifications, triage safety reports, audit immutable system ledger."),
        ("Hotel Food Donor", "donor.taj@annasetu.ai", "Demo@123456", "Post surplus food with AI co-pilot, review proximity match scores, accept requests."),
        ("Verified NGO Shelter", "ngo.hope@annasetu.ai", "Demo@123456", "Natural language AI search, reserve surplus batches within daily shelter capacity."),
        ("Volunteer Courier", "volunteer.rahul@annasetu.ai", "Demo@123456", "Claim unassigned runs, advance delivery stages (ASSIGNED → COMPLETED).")
    ]

    for row_idx, r_data in enumerate(rows9):
        for col_idx, val in enumerate(r_data):
            cell = table9.cell(row_idx + 1, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = c_white if row_idx % 2 == 0 else c_card_bg
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(10)
            p.font.color.rgb = c_dark if col_idx == 0 else c_body
            if col_idx in [0, 2]:
                p.font.bold = True

    # ==========================================
    # SLIDE 10: CONCLUSION & FUTURE ROADMAP
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Conclusion & Future Roadmap")

    add_card(s10, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2),
             "Project Impact & Summary",
             [
                 "Transforms food waste into dignified meals for vulnerable citizens.",
                 "Proves that full-stack web engineering & responsible AI can solve humanitarian logistics bottlenecks.",
                 "Strict adherence to FSSAI standards guarantees trust, hygiene, and accountability.",
                 "Measurable acceleration of UN SDG 2 (Zero Hunger) & SDG 12 (Responsible Consumption).",
                 "Fully functional, demonstrable, and academic presentation ready."
             ], title_color=c_primary)

    add_card(s10, Inches(6.8), Inches(1.7), Inches(5.6), Inches(5.2),
             "Future Engineering Roadmap",
             [
                 "ONDC Logistics Protocol: Open network integration for automated third-party delivery dispatch.",
                 "IoT Cold-Chain Sensors: Bluetooth Low Energy (BLE) temperature beacons logging transit compliance in real time.",
                 "Regional Voice Interface: AI conversational co-pilot in Kannada, Hindi, and Tamil for kitchen staff.",
                 "Predictive Surplus Forecasting: Machine learning models predicting meal surplus before events conclude."
             ], title_color=c_accent)

    prs.save(filename)
    print(f"Presentation PPTX successfully generated at: {filename}")

if __name__ == "__main__":
    create_deck()
